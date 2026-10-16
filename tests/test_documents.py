from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

from documents import DocumentError, MAX_FILE_BYTES, extract_text


def make_pdf(*, blank=False, encrypted=False):
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    if not blank:
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font}),
        })
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 10 100 Td (Biblioteca abre aos sabados.) Tj ET")
        page[NameObject("/Contents")] = stream
    if encrypted:
        writer.encrypt("senha")
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def test_txt_unicode_bom_and_uppercase_extension():
    assert extract_text("estudo.TXT", "\ufeff  Olá, ação!  ".encode()) == "Olá, ação!"


def test_pdf_text():
    assert "Biblioteca abre aos sabados." in extract_text("teste.pdf", make_pdf())


def test_docx_keeps_paragraph_and_table_order():
    document = Document()
    document.add_paragraph("Início: ação")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Prazo"
    table.cell(0, 1).text = "14 dias"
    document.add_paragraph("Fim")
    buffer = BytesIO()
    document.save(buffer)
    assert extract_text("teste.docx", buffer.getvalue()) == "Início: ação\nPrazo | 14 dias\nFim"


@pytest.mark.parametrize(("filename", "content", "message"), [
    ("teste.exe", b"texto", "Formato não suportado"),
    ("teste.txt", b"", "arquivo está vazio"),
    ("teste.txt", b" \n\t", "não contém texto"),
    ("teste.txt", b"\xff\xfe\x00", "UTF-8"),
    ("teste.pdf", b"arquivo quebrado", "Não foi possível ler"),
    ("teste.docx", b"arquivo quebrado", "Não foi possível ler"),
])
def test_invalid_document(filename, content, message):
    with pytest.raises(DocumentError, match=message):
        extract_text(filename, content)


def test_file_too_large():
    with pytest.raises(DocumentError, match="10 MB"):
        extract_text("grande.txt", b"a" * (MAX_FILE_BYTES + 1))


def test_file_at_exact_limit():
    assert len(extract_text("limite.txt", b"a" * 10_000_000)) == 10_000_000


def test_blank_pdf():
    with pytest.raises(DocumentError, match="não contém texto"):
        extract_text("vazio.pdf", make_pdf(blank=True))


def test_encrypted_pdf():
    with pytest.raises(DocumentError, match="protegido"):
        extract_text("protegido.pdf", make_pdf(encrypted=True))


def test_empty_docx():
    buffer = BytesIO()
    Document().save(buffer)
    with pytest.raises(DocumentError, match="não contém texto"):
        extract_text("vazio.docx", buffer.getvalue())
